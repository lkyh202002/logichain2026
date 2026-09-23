"""Build the LogiChain 2025 Base-Case Financial Model (pre price-increase baseline).

Reads the case workbook, keeps its four data sheets untouched and adds formula-driven
model sheets on top of them. Every number in the model is either a link to the data
sheets (green), an explicit input (blue) or a formula (black).

Usage:  python scripts/build_base_case_model.py
"""

from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "LogiChain 2025_Data Case Study Hackathon_ LOGin to LogiChain.xlsx"
OUT = ROOT / "models" / "LogiChain_2025_Base_Case_Financial_Model.xlsx"

# Sheet references used inside formulas
BOM = "'BOM & COGS Ratio'"
SP = "'Supplier Profile'"
FI = "'Factory Information'"
AS = "Assumptions"
BB = "'BOM Cost Build-up'"
RM = "'RM & Supplier Cost'"
CC = "'Capacity Check'"

# ---------------------------------------------------------------- styles
FONT = "Arial"
BLUE, GREEN, BLACK, WHITE = "0000FF", "008000", "000000", "FFFFFF"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HDR_FILL = PatternFill("solid", fgColor="D9E1F2")
SEC_FILL = PatternFill("solid", fgColor="1F3864")
TOT_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="808080")

USD2 = '$#,##0.00;($#,##0.00);"-"'
USD4 = '$#,##0.0000;($#,##0.0000);"-"'
USD0 = '$#,##0;($#,##0);"-"'
PCT1 = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
NUM0 = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
IDX4 = '0.0000;(0.0000);"-"'


def put(ws, ref, value, fmt=None, bold=False, kind=None, fill=None, align=None,
        wrap=False, italic=False, size=10, color=None):
    """Write a cell and colour it by role: input=blue, cross-sheet=green, calc=black."""
    c = ws[ref]
    c.value = value
    if kind is None:
        if isinstance(value, str) and value.startswith("="):
            kind = "link" if "!" in value else "calc"
        elif isinstance(value, (int, float)):
            kind = "input"
        else:
            kind = "label"
    col = color or {"link": GREEN, "calc": BLACK, "input": BLUE, "label": BLACK}[kind]
    c.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=col)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    return c


def title(ws, text, sub):
    put(ws, "A1", text, bold=True, size=14, color="1F3864")
    put(ws, "A2", sub, italic=True, color="595959")
    ws.sheet_view.showGridLines = False


def section(ws, row, text, first="A", last="J"):
    cols = range(openpyxl.utils.column_index_from_string(first),
                 openpyxl.utils.column_index_from_string(last) + 1)
    for i, col in enumerate(cols):
        c = ws.cell(row=row, column=col)
        c.fill = SEC_FILL
        if i == 0:
            c.value = text
        c.font = Font(name=FONT, size=11, bold=True, color=WHITE)


def header(ws, row, labels, start="A", height=42):
    col0 = openpyxl.utils.column_index_from_string(start)
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=lab)
        c.font = Font(name=FONT, size=10, bold=True)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=THIN, top=THIN)
    ws.row_dimensions[row].height = height


def total_row(ws, row, first, last):
    for col in range(openpyxl.utils.column_index_from_string(first),
                     openpyxl.utils.column_index_from_string(last) + 1):
        c = ws.cell(row=row, column=col)
        c.fill = TOT_FILL
        c.border = Border(top=THIN, bottom=THIN)
        rgb = c.font.color.rgb if (c.value is not None and c.font.color is not None) else None
        c.font = Font(name=FONT, size=10, bold=True, color=rgb if isinstance(rgb, str) else BLACK)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def status_format(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(
        operator="equal", formula=['"OK"'],
        font=Font(name=FONT, bold=True, color="006100"),
        fill=PatternFill("solid", fgColor="C6EFCE")))
    for bad in ('"ERROR"', '"Shortfall"'):
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=[bad],
            font=Font(name=FONT, bold=True, color="9C0006"),
            fill=PatternFill("solid", fgColor="FFC7CE")))


def anchor(ws, col, row):
    """Top-left cell of the merged range that contains (col, row)."""
    ref = f"{col}{row}"
    for rng in ws.merged_cells.ranges:
        if ref in rng:
            return f"{openpyxl.utils.get_column_letter(rng.min_col)}{rng.min_row}"
    return ref


# ---------------------------------------------------------------- source map
wb = openpyxl.load_workbook(SRC)
bom_ws, sp_ws, fi_ws = wb["BOM & COGS Ratio"], wb["Supplier Profile"], wb["Factory Information"]

L1_ROWS = [r for r in range(3, bom_ws.max_row + 1) if bom_ws[f"A{r}"].value == 1]
L2_ROWS = [r for r in range(3, bom_ws.max_row + 1) if bom_ws[f"A{r}"].value == 2]
SP_ROWS = [r for r in range(3, 14) if sp_ws[f"C{r}"].value]
FI_ROWS = [r for r in range(3, 20) if fi_ws[f"B{r}"].value]
FI_ROW_OF = {fi_ws[f"B{r}"].value: r for r in FI_ROWS}
assert len(L1_ROWS) == 16 and len(L2_ROWS) == 22 and len(SP_ROWS) == 11

# unique materials (first supplier row of each material) and suppliers, in source order
MAT_ROWS = [r for r in SP_ROWS if anchor(sp_ws, "A", r) == f"A{r}"]
SUP_ROWS, seen = [], set()
for r in SP_ROWS:
    if sp_ws[f"C{r}"].value not in seen:
        seen.add(sp_ws[f"C{r}"].value)
        SUP_ROWS.append(r)
ORG_ROWS, seen = [], set()
for r in SP_ROWS:
    if sp_ws[f"D{r}"].value not in seen:
        seen.add(sp_ws[f"D{r}"].value)
        ORG_ROWS.append(r)
# production lines: anchor row of column A
LINE_ROWS = [r for r in FI_ROWS if fi_ws[f"A{r}"].value]
LINE_PARTS = {r: [x for x in FI_ROWS if anchor(fi_ws, "A", x) == f"A{r}"] for r in LINE_ROWS}

n1, n2, nS, nM = len(L1_ROWS), len(L2_ROWS), len(SP_ROWS), len(MAT_ROWS)
nSup, nOrg = len(SUP_ROWS), len(ORG_ROWS)

# ================================================================ Assumptions
A = wb.create_sheet("Assumptions")
title(A, "Assumptions & Inputs",
      "Base case = supplier prices before any price-increase announcement (all price levers = 0%). "
      "Blue = input · Green = link to another sheet · Black = formula · Yellow = key assumption")

section(A, 4, "A. General", "A", "J")
put(A, "B5", "Currency")
put(A, "C5", "USD", kind="input")
put(A, "E5", "All costs in the case data are quoted in USD", italic=True)
put(A, "B6", "Product item code")
put(A, "C6", f"={BOM}!I3")
put(A, "E6", "Source: BOM & COGS Ratio, Product Item Code", italic=True)
put(A, "B7", "Product name")
put(A, "C7", f"={FI}!C3")
put(A, "E7", "Source: Factory Information, Line L0", italic=True)
put(A, "B8", "Months per year")
put(A, "C8", 12, fmt=NUM0)
put(A, "D8", "months")
put(A, "E8", "Assumption: 12-month run-rate (no seasonality / demand data in the dataset)", italic=True)

section(A, 10, "B. Production Volume", "A", "J")
put(A, "B11", "L0 assembly line - monthly capacity")
put(A, "C11", f"={FI}!E3", fmt=NUM0)
put(A, "D11", "sets/month")
put(A, "E11", "Source: Factory Information, Line L0 (FB_SET_001 Front Brake Assembly)", italic=True)
put(A, "B12", "L0 assembly line - utilization")
put(A, "C12", f"={FI}!G3", fmt=PCT1)
put(A, "D12", "%")
put(A, "E12", "Source: Factory Information (capacity actually utilized vs. maximum capacity)", italic=True)
put(A, "B13", "Calculated monthly output")
put(A, "C13", "=C11*C12", fmt=NUM0)
put(A, "D13", "sets/month")
put(A, "E13", "Capacity x Utilization (OEE not applied: utilization already reflects actual output)", italic=True)
put(A, "B14", "Monthly output override (optional)")
put(A, "C14", None, fmt=NUM0, kind="input", fill=YELLOW)
put(A, "D14", "sets/month")
put(A, "E14", "Leave blank to use calculated output; enter the demand figure here if the case brief gives one", italic=True)
put(A, "B15", "Base-case monthly output (used)", bold=True)
put(A, "C15", "=IF(ISNUMBER(C14),C14,C13)", fmt=NUM0, fill=YELLOW, bold=True)
put(A, "D15", "sets/month")
put(A, "E15", "Key assumption - drives monthly / annual COGS and material requirements", italic=True)
put(A, "B16", "Base-case annual output")
put(A, "C16", "=C15*C8", fmt=NUM0)
put(A, "D16", "sets/year")

section(A, 18, "C. Raw-Material Supplier Prices & Price-Change Levers", "A", "J")
header(A, 19, ["#", "Material Code", "Material Name", "Supplier", "Origin", "Supply Ratio",
               "Base Unit Price (USD/pc)", "Price Change vs Base", "Scenario Unit Price (USD/pc)",
               "Source"])
AS_R0 = 20
for i, r in enumerate(SP_ROWS):
    row = AS_R0 + i
    put(A, f"A{row}", i + 1, kind="label", align="center")
    put(A, f"B{row}", f"={SP}!{anchor(sp_ws, 'A', r)}")
    put(A, f"C{row}", f"={SP}!{anchor(sp_ws, 'B', r)}")
    put(A, f"D{row}", f"={SP}!C{r}")
    put(A, f"E{row}", f"={SP}!D{r}")
    put(A, f"F{row}", f"={SP}!E{r}", fmt=PCT1)
    put(A, f"G{row}", f"={SP}!I{r}", fmt=USD2)
    put(A, f"H{row}", 0, fmt=PCT1, fill=YELLOW)
    put(A, f"I{row}", f"=G{row}*(1+H{row})", fmt=USD2)
    put(A, f"J{row}", f"Supplier Profile row {r}", italic=True)
AS_R1 = AS_R0 + nS - 1
nr = AS_R1 + 1
put(A, f"B{nr}", "Base case: every lever in column H stays at 0%. When a price increase is announced, enter it "
                 "as a fraction (10% = 0.10) for the affected supplier / material; Base-case outputs stay frozen "
                 "and the Scenario columns show the impact.", italic=True, color="595959")

section(A, nr + 2, "D. Sensitivity Steps (price increase applied to one supplier)", "A", "J")
SENS_ROW = nr + 3
put(A, f"B{SENS_ROW}", "Price increase steps")
for j, v in enumerate([0.05, 0.10, 0.15, 0.20]):
    put(A, f"{'CDEF'[j]}{SENS_ROW}", v, fmt=PCT1)
put(A, f"G{SENS_ROW}", "Used in Summary sensitivity table", italic=True)
widths(A, {"A": 5, "B": 36, "C": 22, "D": 22, "E": 14, "F": 11, "G": 14, "H": 13, "I": 14, "J": 22})

AS_CODE = f"{AS}!$B${AS_R0}:$B${AS_R1}"
AS_RATIO = f"{AS}!$F${AS_R0}:$F${AS_R1}"
AS_PRICE = f"{AS}!$G${AS_R0}:$G${AS_R1}"
AS_LEVER = f"{AS}!$H${AS_R0}:$H${AS_R1}"
AS_SCEN = f"{AS}!$I${AS_R0}:$I${AS_R1}"
VOL_M = f"{AS}!$C$15"
VOL_Y = f"{AS}!$C$16"

# ================================================================ BOM Cost Build-up
B = wb.create_sheet("BOM Cost Build-up")
title(B, "BOM Cost Build-up - FB_SET_001 Front Brake Assembly",
      "Source: 'BOM & COGS Ratio' sheet. All values per 1 finished set. COGS Value already includes "
      "Qty per set (source ratio = COGS Value / total).")

B_L1_0 = 6
B_L1_1 = B_L1_0 + n1 - 1
B_L1_T = B_L1_1 + 1
B_L2_0 = B_L1_T + 5
B_L2_1 = B_L2_0 + n2 - 1
B_L2_T = B_L2_1 + 1

section(B, 4, "1. Level 1 - Components (per set)", "A", "Q")
header(B, 5, ["#", "Part Code", "Part Name", "Qty / Set", "Unit", "Production Line",
              "COGS Value (USD/set)", "Unit Cost (USD/pc)", "COGS Share (calc)",
              "COGS Ratio (source)", "|Diff| calc vs source", "Level-2 Roll-up (USD/set)",
              "|Diff| roll-up vs L1", "Line Waste Rate", "Waste-adj. COGS (USD/set) - memo",
              "Scenario COGS (USD/set)", "Δ vs Base (USD/set)"])
L2P = f"$B${B_L2_0}:$B${B_L2_1}"
for i, r in enumerate(L1_ROWS):
    row = B_L1_0 + i
    code = bom_ws[f"B{r}"].value
    fr = FI_ROW_OF[code]
    put(B, f"A{row}", i + 1, kind="label", align="center")
    put(B, f"B{row}", f"={BOM}!B{r}")
    put(B, f"C{row}", f"={BOM}!C{r}")
    put(B, f"D{row}", f"={BOM}!E{r}", fmt=NUM0, align="center")
    put(B, f"E{row}", f"={BOM}!F{r}", align="center")
    put(B, f"F{row}", f"={FI}!{anchor(fi_ws, 'A', fr)}", align="center")
    put(B, f"G{row}", f"={BOM}!K{r}", fmt=USD2)
    put(B, f"H{row}", f"=IF(D{row}=0,0,G{row}/D{row})", fmt=USD2)
    put(B, f"I{row}", f"=IF($G${B_L1_T}=0,0,G{row}/$G${B_L1_T})", fmt=PCT2)
    put(B, f"J{row}", f"={BOM}!J{r}", fmt=PCT2)
    put(B, f"K{row}", f"=ROUND(ABS(I{row}-J{row}),6)", fmt=IDX4)
    put(B, f"L{row}", f"=SUMIF({L2P},B{row},$H${B_L2_0}:$H${B_L2_1})", fmt=USD2)
    put(B, f"M{row}", f"=ROUND(ABS(L{row}-G{row}),6)", fmt=IDX4)
    put(B, f"N{row}", f"={FI}!H{fr}", fmt=PCT1)
    put(B, f"O{row}", f"=G{row}/(1-N{row})", fmt=USD2)
    put(B, f"P{row}", f"=SUMIF({L2P},B{row},$L${B_L2_0}:$L${B_L2_1})", fmt=USD2)
    put(B, f"Q{row}", f"=P{row}-G{row}", fmt=USD2)
t = B_L1_T
put(B, f"B{t}", "Total per set", bold=True)
put(B, f"G{t}", f"=SUM(G{B_L1_0}:G{B_L1_1})", fmt=USD2)
put(B, f"I{t}", f"=SUM(I{B_L1_0}:I{B_L1_1})", fmt=PCT2)
put(B, f"J{t}", f"=SUM(J{B_L1_0}:J{B_L1_1})", fmt=PCT2)
put(B, f"K{t}", f"=MAX(K{B_L1_0}:K{B_L1_1})", fmt=IDX4)
put(B, f"L{t}", f"=SUM(L{B_L1_0}:L{B_L1_1})", fmt=USD2)
put(B, f"M{t}", f"=MAX(M{B_L1_0}:M{B_L1_1})", fmt=IDX4)
put(B, f"O{t}", f"=SUM(O{B_L1_0}:O{B_L1_1})", fmt=USD2)
put(B, f"P{t}", f"=SUM(P{B_L1_0}:P{B_L1_1})", fmt=USD2)
put(B, f"Q{t}", f"=SUM(Q{B_L1_0}:Q{B_L1_1})", fmt=USD2)
total_row(B, t, "A", "Q")
put(B, f"B{t + 1}", "Source ratios are rounded to 4 decimals, so |Diff| up to 0.00005 is expected. "
                    "Waste-adj. COGS = COGS / (1 - line waste rate) is a memo only (not in base COGS).",
    italic=True, color="595959")

section(B, B_L2_0 - 2, "2. Level 2 - Raw-Material Decomposition of each Component (per set)", "A", "Q")
header(B, B_L2_0 - 1, ["#", "Parent Code", "Parent Name", "Material Code", "Material Name",
                       "Share of Parent COGS (source)", "Parent COGS (USD/set)",
                       "Material COGS (USD/set, source)", "Recalc = Share x Parent",
                       "|Diff| recalc vs source", "Material Price Index (Scenario/Base)",
                       "Scenario COGS (USD/set)", "Δ vs Base (USD/set)", "% of Total COGS (base)"])
L1C = f"$B${B_L1_0}:$B${B_L1_1}"
for i, r in enumerate(L2_ROWS):
    row = B_L2_0 + i
    put(B, f"A{row}", i + 1, kind="label", align="center")
    put(B, f"B{row}", f"={BOM}!D{r}")
    put(B, f"C{row}", f"=INDEX($C${B_L1_0}:$C${B_L1_1},MATCH(B{row},{L1C},0))")
    put(B, f"D{row}", f"={BOM}!B{r}")
    put(B, f"E{row}", f"={BOM}!C{r}")
    put(B, f"F{row}", f"={BOM}!J{r}", fmt=PCT2)
    put(B, f"G{row}", f"=INDEX($G${B_L1_0}:$G${B_L1_1},MATCH(B{row},{L1C},0))", fmt=USD2)
    put(B, f"H{row}", f"={BOM}!K{r}", fmt=USD4)
    put(B, f"I{row}", f"=F{row}*G{row}", fmt=USD4)
    put(B, f"J{row}", f"=ROUND(ABS(I{row}-H{row}),6)", fmt=IDX4)
    # RM sheet: material codes in column B, price index in column K (rows set below)
    put(B, f"K{row}", f"=INDEX({RM}!$K$6:$K${5 + nM},MATCH(D{row},{RM}!$B$6:$B${5 + nM},0))", fmt=IDX4)
    put(B, f"L{row}", f"=H{row}*K{row}", fmt=USD4)
    put(B, f"M{row}", f"=L{row}-H{row}", fmt=USD4)
    put(B, f"N{row}", f"=IF($H${B_L2_T}=0,0,H{row}/$H${B_L2_T})", fmt=PCT2)
t = B_L2_T
put(B, f"B{t}", "Total per set", bold=True)
put(B, f"H{t}", f"=SUM(H{B_L2_0}:H{B_L2_1})", fmt=USD4)
put(B, f"I{t}", f"=SUM(I{B_L2_0}:I{B_L2_1})", fmt=USD4)
put(B, f"J{t}", f"=MAX(J{B_L2_0}:J{B_L2_1})", fmt=IDX4)
put(B, f"L{t}", f"=SUM(L{B_L2_0}:L{B_L2_1})", fmt=USD4)
put(B, f"M{t}", f"=SUM(M{B_L2_0}:M{B_L2_1})", fmt=USD4)
put(B, f"N{t}", f"=SUM(N{B_L2_0}:N{B_L2_1})", fmt=PCT2)
total_row(B, t, "A", "Q")
put(B, f"B{t + 1}", "Level 2 splits 100% of each component's COGS into raw materials, so the model treats "
                    "component COGS as fully raw-material driven (no labour/overhead split is given in the data).",
    italic=True, color="595959")
widths(B, {"A": 5, "B": 20, "C": 28, "D": 20, "E": 18, "F": 11, "G": 12, "H": 13, "I": 12,
           "J": 12, "K": 12, "L": 12, "M": 12, "N": 11, "O": 13, "P": 12, "Q": 12})
B.freeze_panes = "D4"

# ================================================================ RM & Supplier Cost
R = wb.create_sheet("RM & Supplier Cost")
title(R, "Raw-Material & Supplier Cost Allocation (per set, monthly, annual)",
      "Multi-sourced materials are split by value share = supply ratio x unit price / Σ(supply ratio x unit price). "
      "Scenario material cost = Base x (weighted scenario price / weighted base price).")

M0, M1 = 6, 5 + nM
MT = M1 + 1
S0 = MT + 4
S1 = S0 + nS - 1
ST = S1 + 1
U0 = ST + 4
U1 = U0 + nSup - 1
UT = U1 + 1
O0 = UT + 4
O1 = O0 + nOrg - 1
OT = O1 + 1

section(R, M0 - 2, "1. Cost by Raw Material (per set)", "A", "P")
header(R, M0 - 1, ["#", "Material Code", "Material Name", "# BOM Lines", "# Suppliers",
                   "Σ Supply Ratio", "Base COGS (USD/set)", "% of Total COGS",
                   "Wtd Avg Base Price (USD/pc)", "Wtd Avg Scenario Price (USD/pc)",
                   "Price Index (Scenario/Base)", "Scenario COGS (USD/set)", "Δ vs Base (USD/set)",
                   "Implied Purchase Qty (pcs/set)", "Monthly Requirement (pcs)"])
L2M = f"{BB}!$D${B_L2_0}:$D${B_L2_1}"
L2V = f"{BB}!$H${B_L2_0}:$H${B_L2_1}"
for i, r in enumerate(MAT_ROWS):
    row = M0 + i
    put(R, f"A{row}", i + 1, kind="label", align="center")
    put(R, f"B{row}", f"={SP}!A{r}")
    put(R, f"C{row}", f"={SP}!B{r}")
    put(R, f"D{row}", f"=COUNTIF({L2M},B{row})", fmt=NUM0, align="center")
    put(R, f"E{row}", f"=COUNTIF({AS_CODE},B{row})", fmt=NUM0, align="center")
    put(R, f"F{row}", f"=SUMIF({AS_CODE},B{row},{AS_RATIO})", fmt=PCT1)
    put(R, f"G{row}", f"=SUMIF({L2M},B{row},{L2V})", fmt=USD2)
    put(R, f"H{row}", f"=IF($G${MT}=0,0,G{row}/$G${MT})", fmt=PCT2)
    put(R, f"I{row}", f"=IF(F{row}=0,0,SUMPRODUCT(({AS_CODE}=B{row})*{AS_RATIO}*{AS_PRICE})/F{row})", fmt=USD2)
    put(R, f"J{row}", f"=IF(F{row}=0,0,SUMPRODUCT(({AS_CODE}=B{row})*{AS_RATIO}*{AS_SCEN})/F{row})", fmt=USD2)
    put(R, f"K{row}", f"=IF(I{row}=0,1,J{row}/I{row})", fmt=IDX4)
    put(R, f"L{row}", f"=G{row}*K{row}", fmt=USD2)
    put(R, f"M{row}", f"=L{row}-G{row}", fmt=USD2)
    put(R, f"N{row}", f"=IF(I{row}=0,0,G{row}/I{row})", fmt=NUM2)
    put(R, f"O{row}", f"=N{row}*{VOL_M}", fmt=NUM0)
put(R, f"B{MT}", "Total", bold=True)
put(R, f"D{MT}", f"=SUM(D{M0}:D{M1})", fmt=NUM0, align="center")
put(R, f"G{MT}", f"=SUM(G{M0}:G{M1})", fmt=USD2)
put(R, f"H{MT}", f"=SUM(H{M0}:H{M1})", fmt=PCT2)
put(R, f"L{MT}", f"=SUM(L{M0}:L{M1})", fmt=USD2)
put(R, f"M{MT}", f"=SUM(M{M0}:M{M1})", fmt=USD2)
total_row(R, MT, "A", "P")
put(R, f"B{MT + 1}", "Implied purchase qty = Base COGS / weighted average unit price (memo: translates the BOM "
                     "cost into pieces bought at Supplier Profile prices).", italic=True, color="595959")

section(R, S0 - 2, "2. Cost by Supplier x Material (per set, monthly, annual)", "A", "P")
header(R, S0 - 1, ["#", "Supplier", "Origin", "Material Code", "Material Name", "Supply Ratio",
                   "Base Unit Price (USD/pc)", "Value Share in Material", "Base COGS (USD/set)",
                   "% of Total COGS", "Base Monthly Spend (USD)", "Base Annual Spend (USD)",
                   "Price Change vs Base", "Scenario COGS (USD/set)", "Δ vs Base (USD/set)",
                   "Implied Monthly Qty (pcs)"])
for i in range(nS):
    row, a = S0 + i, AS_R0 + i
    put(R, f"A{row}", i + 1, kind="label", align="center")
    put(R, f"B{row}", f"={AS}!D{a}")
    put(R, f"C{row}", f"={AS}!E{a}")
    put(R, f"D{row}", f"={AS}!B{a}")
    put(R, f"E{row}", f"={AS}!C{a}")
    put(R, f"F{row}", f"={AS}!F{a}", fmt=PCT1)
    put(R, f"G{row}", f"={AS}!G{a}", fmt=USD2)
    put(R, f"H{row}", f"=IFERROR(F{row}*G{row}/SUMPRODUCT(($D${S0}:$D${S1}=D{row})*$F${S0}:$F${S1}*$G${S0}:$G${S1}),0)",
        fmt=PCT1)
    put(R, f"I{row}", f"=INDEX($G${M0}:$G${M1},MATCH(D{row},$B${M0}:$B${M1},0))*H{row}", fmt=USD2)
    put(R, f"J{row}", f"=IF($I${ST}=0,0,I{row}/$I${ST})", fmt=PCT2)
    put(R, f"K{row}", f"=I{row}*{VOL_M}", fmt=USD0)
    put(R, f"L{row}", f"=I{row}*{VOL_Y}", fmt=USD0)
    put(R, f"M{row}", f"={AS}!H{a}", fmt=PCT1)
    put(R, f"N{row}", f"=I{row}*(1+M{row})", fmt=USD2)
    put(R, f"O{row}", f"=N{row}-I{row}", fmt=USD2)
    put(R, f"P{row}", f"=INDEX($O${M0}:$O${M1},MATCH(D{row},$B${M0}:$B${M1},0))*F{row}", fmt=NUM0)
put(R, f"B{ST}", "Total", bold=True)
for col, fmt in (("I", USD2), ("J", PCT2), ("K", USD0), ("L", USD0), ("N", USD2), ("O", USD2)):
    put(R, f"{col}{ST}", f"=SUM({col}{S0}:{col}{S1})", fmt=fmt)
total_row(R, ST, "A", "P")

section(R, U0 - 2, "3. Cost by Supplier", "A", "P")
header(R, U0 - 1, ["#", "Supplier", "Origin", "# Materials Supplied", "Base COGS (USD/set)",
                   "% of Total COGS", "Base Monthly Spend (USD)", "Base Annual Spend (USD)",
                   "Scenario COGS (USD/set)", "Δ vs Base (USD/set)", "Δ vs Base (%)"])
SUPR = f"$B${S0}:$B${S1}"
for i, r in enumerate(SUP_ROWS):
    row = U0 + i
    put(R, f"A{row}", i + 1, kind="label", align="center")
    put(R, f"B{row}", f"={SP}!C{r}")
    put(R, f"C{row}", f"=INDEX($C${S0}:$C${S1},MATCH(B{row},{SUPR},0))")
    put(R, f"D{row}", f"=COUNTIF({SUPR},B{row})", fmt=NUM0, align="center")
    put(R, f"E{row}", f"=SUMIF({SUPR},B{row},$I${S0}:$I${S1})", fmt=USD2)
    put(R, f"F{row}", f"=IF($E${UT}=0,0,E{row}/$E${UT})", fmt=PCT1)
    put(R, f"G{row}", f"=SUMIF({SUPR},B{row},$K${S0}:$K${S1})", fmt=USD0)
    put(R, f"H{row}", f"=SUMIF({SUPR},B{row},$L${S0}:$L${S1})", fmt=USD0)
    put(R, f"I{row}", f"=SUMIF({SUPR},B{row},$N${S0}:$N${S1})", fmt=USD2)
    put(R, f"J{row}", f"=I{row}-E{row}", fmt=USD2)
    put(R, f"K{row}", f"=IF(E{row}=0,0,J{row}/E{row})", fmt=PCT1)
put(R, f"B{UT}", "Total", bold=True)
for col, fmt in (("D", NUM0), ("E", USD2), ("F", PCT1), ("G", USD0), ("H", USD0), ("I", USD2), ("J", USD2)):
    put(R, f"{col}{UT}", f"=SUM({col}{U0}:{col}{U1})", fmt=fmt, align="center" if col == "D" else None)
put(R, f"K{UT}", f"=IF(E{UT}=0,0,J{UT}/E{UT})", fmt=PCT1)
total_row(R, UT, "A", "P")

section(R, O0 - 2, "4. Cost by Origin Country", "A", "P")
header(R, O0 - 1, ["#", "Origin", "# Suppliers", "", "Base COGS (USD/set)", "% of Total COGS",
                   "Base Monthly Spend (USD)", "Base Annual Spend (USD)", "Scenario COGS (USD/set)",
                   "Δ vs Base (USD/set)", "Δ vs Base (%)"])
ORGR = f"$C${S0}:$C${S1}"
for i, r in enumerate(ORG_ROWS):
    row = O0 + i
    put(R, f"A{row}", i + 1, kind="label", align="center")
    put(R, f"B{row}", f"={SP}!D{r}")
    put(R, f"C{row}", f"=COUNTIF($C${U0}:$C${U1},B{row})", fmt=NUM0, align="center")
    put(R, f"E{row}", f"=SUMIF({ORGR},B{row},$I${S0}:$I${S1})", fmt=USD2)
    put(R, f"F{row}", f"=IF($E${OT}=0,0,E{row}/$E${OT})", fmt=PCT1)
    put(R, f"G{row}", f"=SUMIF({ORGR},B{row},$K${S0}:$K${S1})", fmt=USD0)
    put(R, f"H{row}", f"=SUMIF({ORGR},B{row},$L${S0}:$L${S1})", fmt=USD0)
    put(R, f"I{row}", f"=SUMIF({ORGR},B{row},$N${S0}:$N${S1})", fmt=USD2)
    put(R, f"J{row}", f"=I{row}-E{row}", fmt=USD2)
    put(R, f"K{row}", f"=IF(E{row}=0,0,J{row}/E{row})", fmt=PCT1)
put(R, f"B{OT}", "Total", bold=True)
for col, fmt in (("C", NUM0), ("E", USD2), ("F", PCT1), ("G", USD0), ("H", USD0), ("I", USD2), ("J", USD2)):
    put(R, f"{col}{OT}", f"=SUM({col}{O0}:{col}{O1})", fmt=fmt, align="center" if col == "C" else None)
put(R, f"K{OT}", f"=IF(E{OT}=0,0,J{OT}/E{OT})", fmt=PCT1)
total_row(R, OT, "A", "P")
widths(R, {"A": 5, "B": 24, "C": 20, "D": 20, "E": 18, "F": 11, "G": 13, "H": 13, "I": 13,
           "J": 13, "K": 14, "L": 14, "M": 12, "N": 13, "O": 13, "P": 13})

# ================================================================ Capacity Check
C = wb.create_sheet("Capacity Check")
title(C, "Capacity Check (memo) - does the base-case volume fit each production line?",
      "Required output = Qty per set x base-case monthly output. Effective output = Monthly capacity x Utilization. "
      "Each line runs 24 hours (Factory Information note).")
header(C, 4, ["Line", "Part Code(s)", "Part Name(s)", "Qty per Set on Line", "Required Output (pcs/month)",
              "Monthly Capacity (pcs)", "Utilization", "Effective Output (pcs/month)", "OEE (info)",
              "Coverage (Effective / Required)", "Max Sets Supported / month", "Status"])
C0 = 5
BBC = f"{BB}!$B${B_L1_0}:$B${B_L1_1}"
BBQ = f"{BB}!$D${B_L1_0}:$D${B_L1_1}"
for i, lr in enumerate(LINE_ROWS):
    row = C0 + i
    parts = LINE_PARTS[lr]
    put(C, f"A{row}", f"={FI}!A{lr}", align="center")
    put(C, f"B{row}", "=" + '&" / "&'.join(f"{FI}!B{p}" for p in parts))
    put(C, f"C{row}", "=" + '&" / "&'.join(f"{FI}!C{p}" for p in parts))
    if fi_ws[f"B{lr}"].value == bom_ws["I3"].value:
        put(C, f"D{row}", 1, fmt=NUM0, align="center")  # the finished set itself
    else:
        put(C, f"D{row}", "=" + "+".join(f"SUMIF({BBC},{FI}!B{p},{BBQ})" for p in parts),
            fmt=NUM0, align="center")
    put(C, f"E{row}", f"=D{row}*{VOL_M}", fmt=NUM0)
    put(C, f"F{row}", f"={FI}!E{lr}", fmt=NUM0)
    put(C, f"G{row}", f"={FI}!G{lr}", fmt=PCT1)
    put(C, f"H{row}", f"=F{row}*G{row}", fmt=NUM0)
    put(C, f"I{row}", f"={FI}!F{lr}", fmt=PCT1)
    put(C, f"J{row}", f"=IF(E{row}=0,0,H{row}/E{row})", fmt=PCT1)
    put(C, f"K{row}", f"=IF(D{row}=0,0,H{row}/D{row})", fmt=NUM0)
    put(C, f"L{row}", f'=IF(J{row}>=1,"OK","Shortfall")', align="center")
C1 = C0 + len(LINE_ROWS) - 1
status_format(C, f"L{C0}:L{C1}")
put(C, f"B{C1 + 2}", "Bottleneck - max sets/month supported", bold=True)
put(C, f"E{C1 + 2}", f"=MIN(K{C0}:K{C1})", fmt=NUM0, bold=True)
put(C, f"B{C1 + 3}", "Bottleneck line")
put(C, f"E{C1 + 3}", f"=INDEX(A{C0}:A{C1},MATCH(E{C1 + 2},K{C0}:K{C1},0))", align="right")
put(C, f"B{C1 + 4}", "# lines with shortfall at base volume")
put(C, f"E{C1 + 4}", f'=COUNTIF(L{C0}:L{C1},"Shortfall")', fmt=NUM0)
put(C, f"B{C1 + 6}", "Reading: capacity is treated as pieces per month. If several component lines show a "
                     "shortfall, check the case brief - capacity may be quoted in set-equivalents, or some parts "
                     "may be bought in. Per-set COGS is not affected by volume; only monthly / annual totals are.",
    italic=True, color="595959")
widths(C, {"A": 7, "B": 20, "C": 42, "D": 11, "E": 14, "F": 13, "G": 11, "H": 14, "I": 10,
           "J": 14, "K": 14, "L": 12})
CAP_BOTTLENECK = f"{CC}!$E${C1 + 2}"
CAP_LINE = f"{CC}!$E${C1 + 3}"
CAP_SHORT = f"{CC}!$E${C1 + 4}"

# ================================================================ Summary
S = wb.create_sheet("Summary")
title(S, "Base-Case Summary - FB_SET_001 Front Brake Assembly",
      "Baseline before the price-increase announcement · USD · per set and monthly / annual run-rate")
put(S, "B3", f'=IF(SUMPRODUCT(ABS({AS_LEVER}))=0,"Status: BASE CASE - all price levers = 0% (Scenario = Base)",'
             f'"Status: PRICE LEVERS ACTIVE - Scenario column differs from Base")', bold=True, color="C00000")

section(S, 5, "1. Key Metrics", "A", "K")
header(S, 6, ["", "Metric", "Base Case", "Scenario (price levers)", "Δ vs Base", "Δ %"], height=30)
KM = [("COGS per set (USD)", f"={BB}!G{B_L1_T}", f"={BB}!P{B_L1_T}", USD2),
      ("Monthly output (sets)", f"={VOL_M}", f"={VOL_M}", NUM0),
      ("Annual output (sets)", f"={VOL_Y}", f"={VOL_Y}", NUM0),
      ("Monthly COGS (USD)", "=C7*C8", "=D7*D8", USD0),
      ("Annual COGS (USD)", "=C7*C9", "=D7*D9", USD0)]
for i, (lab, base, scen, fmt) in enumerate(KM):
    row = 7 + i
    put(S, f"B{row}", lab, bold=(i == 0))
    put(S, f"C{row}", base, fmt=fmt, bold=(i == 0))
    put(S, f"D{row}", scen, fmt=fmt, bold=(i == 0))
    put(S, f"E{row}", f"=D{row}-C{row}", fmt=fmt)
    put(S, f"F{row}", f"=IF(C{row}=0,0,E{row}/C{row})", fmt=PCT1)
put(S, "B12", "Model checks")
put(S, "C12", "=Checks!$C$3", bold=True)
status_format(S, "C12")


def cost_table(ws, top, heading, cols, n, total_src):
    """Stack a cost-structure table that links row-by-row to a calc table."""
    section(ws, top, heading, "A", "K")
    header(ws, top + 1, ["#"] + [c[0] for c in cols])
    r0 = top + 2
    r1 = r0 + n - 1
    for i in range(n):
        row = r0 + i
        put(ws, f"A{row}", i + 1, kind="label", align="center")
        for j, (_, f, fmt) in enumerate(cols):
            col = "BCDEFGHIJK"[j]
            if f is None:
                continue
            val = f(row, r0, r1)
            put(ws, f"{col}{row}", val, fmt=fmt, align="center" if fmt == NUM0 and col in "DG" else None)
    t = r1 + 1
    put(ws, f"B{t}", "Total", bold=True)
    for j, (_, f, fmt) in enumerate(cols):
        col = "BCDEFGHIJK"[j]
        if col in total_src:
            put(ws, f"{col}{t}", total_src[col](r0, r1), fmt=fmt)
    total_row(ws, t, "A", "K")
    return r0, r1, t


def rank(col):
    return lambda row, r0, r1: f"=RANK({col}{row},${col}${r0}:${col}${r1})"


def share(col):
    return lambda row, r0, r1: f"=IF(${col}${r1 + 1}=0,0,{col}{row}/${col}${r1 + 1})"


def sums(*cols):
    return {c: (lambda r0, r1, c=c: f"=SUM({c}{r0}:{c}{r1})") for c in cols}


# 2. by component
top = 14
comp_r0, comp_r1, comp_t = cost_table(
    S, top, "2. COGS Structure by Component (Level 1)",
    [("Part Code", None, None), ("Part Name", None, None), ("Qty / Set", None, NUM0),
     ("COGS (USD/set)", None, USD2), ("% of COGS", share("E"), PCT1), ("Rank", rank("E"), NUM0),
     ("Annual COGS (USD)", lambda row, r0, r1: f"=E{row}*{VOL_Y}", USD0),
     ("Scenario (USD/set)", None, USD2),
     ("Δ vs Base (USD/set)", lambda row, r0, r1: f"=I{row}-E{row}", USD2)],
    n1, sums("E", "F", "H", "I", "J"))
for i in range(n1):
    row, src = comp_r0 + i, B_L1_0 + i
    for col, scol, fmt in (("B", "B", None), ("C", "C", None), ("D", "D", NUM0), ("E", "G", USD2), ("I", "P", USD2)):
        put(S, f"{col}{row}", f"={BB}!{scol}{src}", fmt=fmt, align="center" if col == "D" else None)

# 3. by raw material
top = comp_t + 2
mat_r0, mat_r1, mat_t = cost_table(
    S, top, "3. COGS Structure by Raw Material (Level 2)",
    [("Material Code", None, None), ("Material Name", None, None), ("# Suppliers", None, NUM0),
     ("COGS (USD/set)", None, USD2), ("% of COGS", share("E"), PCT1), ("Rank", rank("E"), NUM0),
     ("Annual COGS (USD)", lambda row, r0, r1: f"=E{row}*{VOL_Y}", USD0),
     ("Scenario (USD/set)", None, USD2),
     ("Δ vs Base (USD/set)", lambda row, r0, r1: f"=I{row}-E{row}", USD2)],
    nM, sums("E", "F", "H", "I", "J"))
for i in range(nM):
    row, src = mat_r0 + i, M0 + i
    for col, scol, fmt in (("B", "B", None), ("C", "C", None), ("D", "E", NUM0), ("E", "G", USD2), ("I", "L", USD2)):
        put(S, f"{col}{row}", f"={RM}!{scol}{src}", fmt=fmt, align="center" if col == "D" else None)

# 4. by supplier
top = mat_t + 2
sup_r0, sup_r1, sup_t = cost_table(
    S, top, "4. COGS by Supplier",
    [("Supplier", None, None), ("Origin", None, None), ("# Materials", None, NUM0),
     ("COGS (USD/set)", None, USD2), ("% of COGS", share("E"), PCT1), ("Rank", rank("E"), NUM0),
     ("Annual Spend (USD)", lambda row, r0, r1: f"=E{row}*{VOL_Y}", USD0),
     ("Scenario (USD/set)", None, USD2),
     ("Δ vs Base (USD/set)", lambda row, r0, r1: f"=I{row}-E{row}", USD2)],
    nSup, sums("E", "F", "H", "I", "J"))
for i in range(nSup):
    row, src = sup_r0 + i, U0 + i
    for col, scol, fmt in (("B", "B", None), ("C", "C", None), ("D", "D", NUM0), ("E", "E", USD2), ("I", "I", USD2)):
        put(S, f"{col}{row}", f"={RM}!{scol}{src}", fmt=fmt, align="center" if col == "D" else None)

# 5. by origin
top = sup_t + 2
org_r0, org_r1, org_t = cost_table(
    S, top, "5. COGS by Origin Country",
    [("Origin", None, None), ("# Suppliers", None, NUM0), ("", None, None),
     ("COGS (USD/set)", None, USD2), ("% of COGS", share("E"), PCT1), ("Rank", rank("E"), NUM0),
     ("Annual Spend (USD)", lambda row, r0, r1: f"=E{row}*{VOL_Y}", USD0),
     ("Scenario (USD/set)", None, USD2),
     ("Δ vs Base (USD/set)", lambda row, r0, r1: f"=I{row}-E{row}", USD2)],
    nOrg, sums("E", "F", "H", "I", "J"))
for i in range(nOrg):
    row, src = org_r0 + i, O0 + i
    for col, scol, fmt in (("B", "B", None), ("C", "C", NUM0), ("E", "E", USD2), ("I", "I", USD2)):
        put(S, f"{col}{row}", f"={RM}!{scol}{src}", fmt=fmt, align="center" if col == "C" else None)
put(S, f"C{org_t}", f"=SUM(C{org_r0}:C{org_r1})", fmt=NUM0, bold=True, align="center", fill=TOT_FILL)

# 6. sensitivity
top = org_t + 2
section(S, top, "6. Price Sensitivity - Δ COGS if ONE supplier raises prices on all its materials", "A", "K")
put(S, f"D{top + 1}", "Δ COGS per set (USD)", bold=True, align="center", fill=HDR_FILL)
put(S, f"H{top + 1}", "Δ annual COGS (USD)", bold=True, align="center", fill=HDR_FILL)
S.merge_cells(f"D{top + 1}:G{top + 1}")
S.merge_cells(f"H{top + 1}:K{top + 1}")
header(S, top + 2, ["#", "Supplier", "Base COGS (USD/set)"], height=30)
for j in range(4):
    for col in ("DEFG"[j], "HIJK"[j]):
        c = put(S, f"{col}{top + 2}", f"={AS}!{'CDEF'[j]}{SENS_ROW}", fmt='"+"0%', bold=True, align="center",
                fill=HDR_FILL)
        c.border = Border(top=THIN, bottom=THIN)
sen_r0 = top + 3
for i in range(nSup):
    row = sen_r0 + i
    put(S, f"A{row}", i + 1, kind="label", align="center")
    put(S, f"B{row}", f"=B{sup_r0 + i}")
    put(S, f"C{row}", f"=E{sup_r0 + i}", fmt=USD2)
    for j in range(4):
        dc, ac = "DEFG"[j], "HIJK"[j]
        put(S, f"{dc}{row}", f"=$C{row}*{dc}${top + 2}", fmt=USD2)
        put(S, f"{ac}{row}", f"=$C{row}*{ac}${top + 2}*{VOL_Y}", fmt=USD0)
sen_r1 = sen_r0 + nSup - 1
t = sen_r1 + 1
put(S, f"B{t}", "All suppliers (uniform increase)", bold=True)
put(S, f"C{t}", f"=SUM(C{sen_r0}:C{sen_r1})", fmt=USD2)
for j in range(4):
    for col, fmt in (("DEFG"[j], USD2), ("HIJK"[j], USD0)):
        put(S, f"{col}{t}", f"=SUM({col}{sen_r0}:{col}{sen_r1})", fmt=fmt)
total_row(S, t, "A", "K")
put(S, f"B{t + 1}", "Δ% of total COGS = supplier's share of COGS x price increase. Linear because the model "
                    "passes material price changes 1:1 into COGS.", italic=True, color="595959")

# 7. waste memo
top = t + 3
section(S, top, "7. Memo - Production Waste (NOT included in base COGS)", "A", "K")
W = top + 1
put(S, f"B{W}", "BOM standard COGS per set (USD)")
put(S, f"C{W}", f"={BB}!G{B_L1_T}", fmt=USD2)
put(S, f"B{W + 1}", "Component-level waste-adjusted COGS (USD/set)")
put(S, f"C{W + 1}", f"={BB}!O{B_L1_T}", fmt=USD2)
put(S, f"D{W + 1}", "Σ component COGS / (1 - line waste rate)", italic=True, color="595959")
put(S, f"B{W + 2}", "L0 assembly waste rate")
put(S, f"C{W + 2}", f"={FI}!H3", fmt=PCT1)
put(S, f"B{W + 3}", "Fully waste-adjusted COGS per set (USD)", bold=True)
put(S, f"C{W + 3}", f"=C{W + 1}/(1-C{W + 2})", fmt=USD2, bold=True)
put(S, f"B{W + 4}", "Waste uplift vs BOM (USD/set)")
put(S, f"C{W + 4}", f"=C{W + 3}-C{W}", fmt=USD2)
put(S, f"D{W + 4}", f"=IF(C{W}=0,0,C{W + 4}/C{W})", fmt=PCT1)
put(S, f"B{W + 5}", "Monthly cost of waste at base volume (USD)")
put(S, f"C{W + 5}", f"=C{W + 4}*{VOL_M}", fmt=USD0)
put(S, f"B{W + 6}", "Annual cost of waste (USD)")
put(S, f"C{W + 6}", f"=C{W + 4}*{VOL_Y}", fmt=USD0)
put(S, f"B{W + 7}", "Use only if the case treats the BOM COGS as a standard (zero-waste) cost.",
    italic=True, color="595959")

# 8. capacity memo
top = W + 9
section(S, top, "8. Memo - Capacity Check at Base-Case Volume", "A", "K")
put(S, f"B{top + 1}", "Lines with shortfall")
put(S, f"C{top + 1}", f"={CAP_SHORT}", fmt=NUM0)
put(S, f"B{top + 2}", "Bottleneck line")
put(S, f"C{top + 2}", f"={CAP_LINE}", align="right")
put(S, f"B{top + 3}", "Max sets/month supported by bottleneck")
put(S, f"C{top + 3}", f"={CAP_BOTTLENECK}", fmt=NUM0)
put(S, f"D{top + 3}", "See 'Capacity Check' sheet - verify capacity units against the case brief",
    italic=True, color="595959")
widths(S, {"A": 5, "B": 44, "C": 24, "D": 16, "E": 14, "F": 11, "G": 9, "H": 16, "I": 14, "J": 14, "K": 14})

SUM_COMP = (comp_r0, comp_r1, comp_t)
SUM_MAT = (mat_r0, mat_r1, mat_t)
SUM_SUP = (sup_r0, sup_r1, sup_t)
SUM_ORG = (org_r0, org_r1, org_t)

# ================================================================ Checks
K = wb.create_sheet("Checks")
title(K, "Model Integrity Checks", "Every check must read OK before the baseline is used")
put(K, "B3", "Overall status", bold=True)
header(K, 5, ["#", "Check", "Value", "Target", "Tolerance", "Result"], height=30)
checks = [
    ("Level-1 calculated shares sum to 100%", f"={BB}!I{B_L1_T}", 1, 0.000001),
    ("Level-1 calc share vs source COGS Ratio (max |diff|, source rounded to 4 dp)",
     f"={BB}!K{B_L1_T}", 0, 0.0001),
    ("Level-2 COGS Value = Share x Parent COGS (max |diff|)", f"={BB}!J{B_L2_T}", 0, 0.001),
    ("Each component fully decomposed into raw materials (max |L2 roll-up - L1|)",
     f"={BB}!M{B_L1_T}", 0, 0.001),
    ("Σ Level-2 COGS = Σ Level-1 COGS (USD/set)", f"={BB}!H{B_L2_T}", f"={BB}!G{B_L1_T}", 0.001),
    ("Supply ratios sum to 100% for every material (Σ |Σratio - 1|)",
     f"=SUMPRODUCT(ABS({RM}!F{M0}:F{M1}-1))", 0, 0.000001),
    ("Raw-material roll-up = total COGS (USD/set)", f"={RM}!G{MT}", f"={BB}!G{B_L1_T}", 0.001),
    ("Supplier x material roll-up = total COGS (USD/set)", f"={RM}!I{ST}", f"={BB}!G{B_L1_T}", 0.001),
    ("Supplier roll-up = total COGS (USD/set)", f"={RM}!E{UT}", f"={BB}!G{B_L1_T}", 0.001),
    ("Origin roll-up = total COGS (USD/set)", f"={RM}!E{OT}", f"={BB}!G{B_L1_T}", 0.001),
    ("Scenario: material-based total = supplier-based total (USD/set)", f"={RM}!L{MT}", f"={RM}!N{ST}", 0.001),
    ("Summary COGS per set = SUM of source COGS Value (Level 1)",
     "=Summary!C7", f"=SUM({BOM}!K{L1_ROWS[0]}:K{L1_ROWS[-1]})", 0.001),
]
k0 = 6
for i, (lab, val, tgt, tol) in enumerate(checks):
    row = k0 + i
    put(K, f"A{row}", i + 1, kind="label", align="center")
    put(K, f"B{row}", lab)
    put(K, f"C{row}", val, fmt="#,##0.000000")
    put(K, f"D{row}", tgt, fmt="#,##0.000000")
    put(K, f"E{row}", tol, fmt="0.000000")
    put(K, f"F{row}", f'=IF(ABS(C{row}-D{row})<=E{row},"OK","ERROR")', align="center")
k1 = k0 + len(checks) - 1
put(K, "C3", f'=IF(COUNTIF(F{k0}:F{k1},"ERROR")=0,"OK","ERROR")', bold=True, align="center")
status_format(K, f"F{k0}:F{k1}")
status_format(K, "C3")
widths(K, {"A": 5, "B": 78, "C": 16, "D": 16, "E": 12, "F": 10})

# ================================================================ Cover
V = wb.create_sheet("Cover")
V.sheet_view.showGridLines = False
put(V, "B2", "LogiChain 2025 - Base-Case Financial Model", bold=True, size=16, color="1F3864")
put(V, "B3", "Baseline COGS của FB_SET_001 (Front Brake Assembly) TRƯỚC khi có thông báo tăng giá",
    italic=True, size=11, color="595959")
put(V, "B4", "Version v1.0 - Base case · Nguồn: LogiChain 2025 Data Case Study Hackathon (LOGin to LogiChain)",
    italic=True, color="595959")

section(V, 6, "Kết quả chính (Base case)", "B", "E")
kr = [("COGS mỗi bộ (USD/set)", "=Summary!C7", USD2),
      ("Sản lượng base (bộ/tháng)", "=Summary!C8", NUM0),
      ("COGS tháng (USD)", "=Summary!C10", USD0),
      ("COGS năm (USD)", "=Summary!C11", USD0),
      ("NVL lớn nhất - % COGS",
       f"=INDEX(Summary!C{mat_r0}:C{mat_r1},MATCH(1,Summary!G{mat_r0}:G{mat_r1},0))&\" - \""
       f"&TEXT(MAX(Summary!F{mat_r0}:F{mat_r1}),\"0.0%\")", None),
      ("NCC lớn nhất - % COGS",
       f"=INDEX(Summary!B{sup_r0}:B{sup_r1},MATCH(1,Summary!G{sup_r0}:G{sup_r1},0))&\" - \""
       f"&TEXT(MAX(Summary!F{sup_r0}:F{sup_r1}),\"0.0%\")", None),
      ("Quốc gia phụ thuộc nhiều nhất - % COGS",
       f"=INDEX(Summary!B{org_r0}:B{org_r1},MATCH(1,Summary!G{org_r0}:G{org_r1},0))&\" - \""
       f"&TEXT(MAX(Summary!F{org_r0}:F{org_r1}),\"0.0%\")", None),
      ("Trạng thái kiểm tra mô hình", "=Checks!C3", None)]
for i, (lab, f, fmt) in enumerate(kr):
    put(V, f"B{7 + i}", lab)
    put(V, f"C{7 + i}", f, fmt=fmt, bold=True, align="right")
status_format(V, f"C{7 + len(kr) - 1}")

r = 7 + len(kr) + 1
section(V, r, "Cấu trúc file", "B", "E")
sheets = [("Summary", "Kết quả: COGS/bộ, COGS tháng/năm, cơ cấu theo linh kiện - NVL - NCC - quốc gia, "
                      "độ nhạy giá, memo hao hụt & công suất"),
          ("Assumptions", "Mọi giả định & đòn bẩy giá (price levers = 0% ở base case), sản lượng, bước độ nhạy"),
          ("BOM Cost Build-up", "Roll-up chi phí Level 1 (linh kiện) và Level 2 (NVL) từ sheet BOM & COGS Ratio"),
          ("RM & Supplier Cost", "Phân bổ COGS theo NVL, NCC (supply ratio x đơn giá) và quốc gia xuất xứ"),
          ("Capacity Check", "Memo: sản lượng base so với công suất hiệu dụng từng line"),
          ("Checks", "12 kiểm tra tính nhất quán dữ liệu & mô hình"),
          ("BOM & COGS Ratio", "Dữ liệu gốc (giữ nguyên) - cùng 3 sheet dữ liệu gốc còn lại")]
for i, (name, desc) in enumerate(sheets):
    c = put(V, f"B{r + 1 + i}", f'=HYPERLINK("#\'{name}\'!A1","{name}")', bold=True, color="0563C1")
    c.font = Font(name=FONT, size=10, bold=True, color="0563C1", underline="single")
    put(V, f"C{r + 1 + i}", desc)
r = r + 1 + len(sheets) + 1

section(V, r, "Quy ước màu", "B", "E")
legend = [("1,234", BLUE, None, "Chữ xanh dương = input nhập tay / đòn bẩy kịch bản"),
          ("1,234", GREEN, None, "Chữ xanh lá = công thức lấy dữ liệu từ sheet khác"),
          ("1,234", BLACK, None, "Chữ đen = công thức tính trong cùng sheet"),
          ("1,234", BLUE, YELLOW, "Nền vàng = giả định chính / ô cần điền khi có thông tin mới")]
for i, (txt, col, fill, desc) in enumerate(legend):
    put(V, f"B{r + 1 + i}", txt, color=col, fill=fill, align="center", kind="label")
    put(V, f"C{r + 1 + i}", desc)
r = r + 1 + len(legend) + 1

section(V, r, "Phương pháp & giả định chính", "B", "E")
notes = [
    "1. COGS/bộ = Σ COGS Value (USD) của 16 linh kiện Level 1. COGS Value đã bao gồm số lượng/bộ "
    "(kiểm chứng: COGS Ratio của đề = COGS Value / tổng).",
    "2. Level 2 phân rã 100% COGS của linh kiện cha thành NVL (COGS NVL = Ratio x COGS cha) => toàn bộ COGS "
    "được coi là do NVL quyết định (đề không tách nhân công / overhead).",
    "3. NVL nhiều nhà cung cấp: phân bổ giá trị theo value share = supply ratio x đơn giá / Σ(supply ratio x đơn giá).",
    "4. Kịch bản giá: COGS NVL mới = COGS base x (giá bình quân gia quyền mới / giá bình quân gia quyền base). "
    "Base case: mọi price change = 0% => Scenario = Base.",
    "5. Sản lượng base = Monthly capacity L0 x Utilization L0 = 24,000 x 85% = 20,400 bộ/tháng; 12 tháng/năm. "
    "Có ô override nếu đề cho nhu cầu thực tế.",
    "6. Tỷ lệ hao hụt (waste) KHÔNG đưa vào base COGS (giữ đúng số liệu BOM); chỉ trình bày ở dạng memo.",
    "7. OEE không dùng để tính sản lượng vì Utilization (theo định nghĩa của đề) đã phản ánh mức sử dụng thực tế.",
]
for i, n in enumerate(notes):
    put(V, f"B{r + 1 + i}", n, wrap=True)
    V.merge_cells(f"B{r + 1 + i}:E{r + 1 + i}")
    V.row_dimensions[r + 1 + i].height = 30
r = r + 1 + len(notes) + 1

section(V, r, "Khi có thông báo tăng giá", "B", "E")
howto = [
    "Mở sheet Assumptions, mục C: nhập % tăng giá vào cột H (Price Change vs Base) cho đúng NCC / NVL "
    "(10% = 0.10). Nếu đề cho giá mới dạng USD/pc: % = giá mới / giá cũ - 1.",
    "Summary tự tính cột Scenario và Δ; cột Base Case giữ nguyên làm mốc so sánh. "
    "Bảng độ nhạy (mục 6 của Summary) cho biết trước mức tác động khi từng NCC tăng giá 5-20%.",
]
for i, n in enumerate(howto):
    put(V, f"B{r + 1 + i}", n, wrap=True)
    V.merge_cells(f"B{r + 1 + i}:E{r + 1 + i}")
    V.row_dimensions[r + 1 + i].height = 30
widths(V, {"A": 3, "B": 40, "C": 28, "D": 40, "E": 40})

# ---------------------------------------------------------------- order, tabs, save
new_order = ["Cover", "Summary", "Assumptions", "BOM Cost Build-up", "RM & Supplier Cost",
             "Capacity Check", "Checks", "Terms Explanation", "BOM & COGS Ratio",
             "Supplier Profile", "Factory Information"]
wb._sheets = [wb[n] for n in new_order]
tabs = {"Cover": "1F3864", "Summary": "1F3864", "Assumptions": "FFC000", "BOM Cost Build-up": "4472C4",
        "RM & Supplier Cost": "4472C4", "Capacity Check": "4472C4", "Checks": "70AD47"}
for n in new_order:
    wb[n].sheet_properties.tabColor = tabs.get(n, "A6A6A6")
wb.active = 0
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = ws.title == "Cover"

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f"saved {OUT}")
